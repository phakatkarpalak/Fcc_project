"""Build GreenWarm's 24-row India carbon curve from the Grid-India generation workbooks.

Input: the monthly .xlsx files of "Electricity Demand, Solar and Wind Generation Data of India at
1-hour interval" (Mukherjee, Kalita, Kumar, Mendeley Data, doi:10.17632/y58jknpgs8). Sheet1 has,
in row 2, the all-India signals  TOTAL_THM_ONLY|P (thermal), TOTAL_NLDC_HYD|P (hydro),
TOTAL_NLDC_GAS|P (gas), TOTAL_NLDC_NPC|P (nuclear), ALL_INDIA_WIND|P, ALL_IND_SOLAR|P  in MW,
with samples every 5 minutes (2022 files) or 10 seconds (2023 files).

Cleaning: solar below zero is clipped to 0; other negative readings, zero coal and implausibly low totals are dropped.
Method: for every sample, emissions = sum(MW_i x factor_i); then for each hour of the day
intensity = total emissions / total generation over every sample in that hour, computed separately
for each file and then averaged across files, so every file counts equally whatever its sampling
rate. (A generation-weighted average within a file, not an average of ratios.)

Emission factors (gCO2eq/kWh), IPCC AR5 Annex III medians:
  lifecycle (default): coal 820, gas 490, hydro 24, nuclear 12, wind onshore 11, solar PV utility 48
  direct (--basis direct): coal 760, gas 370, hydro 0, nuclear 0, wind 0, solar 0
Override any one with --ef coal=950. The thermal signal is treated as coal/lignite.

Limits you should state in a report: all-India (not a regional) grid; only these six sources, so
small hydro, biomass, rooftop solar and other supply are missing; the factors are generic, not
Indian plant data. Absolute levels are therefore approximate; the DAILY SHAPE is the useful part
(the scheduler decides on intensity relative to its own daily mean).

Usage:
    python prepare_india_carbon.py "July 2022.xlsx" "July 2023.xlsx" india_july.csv
    python prepare_india_carbon.py "July 2022.xlsx" "July 2023.xlsx" india_july_utc.csv --shift-hours -5.5
Timestamps are assumed to be Indian Standard Time, as exported. Use --shift-hours -5.5 for UTC.
"""
import argparse
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

SIGNALS = {  # signal name in row 2  ->  source key
    "TOTAL_THM_ONLY|P": "coal", "TOTAL_NLDC_HYD|P": "hydro", "TOTAL_NLDC_GAS|P": "gas",
    "TOTAL_NLDC_NPC|P": "nuclear", "ALL_INDIA_WIND|P": "wind", "ALL_IND_SOLAR|P": "solar",
}
FACTORS = {
    "lifecycle": {"coal": 820.0, "gas": 490.0, "hydro": 24.0, "nuclear": 12.0, "wind": 11.0, "solar": 48.0},
    "direct": {"coal": 760.0, "gas": 370.0, "hydro": 0.0, "nuclear": 0.0, "wind": 0.0, "solar": 0.0},
}


def read_workbook(path: str) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = None
    for name in wb.sheetnames:
        top = list(wb[name].iter_rows(min_row=2, max_row=2, values_only=True))
        if top and any(str(c).strip() in SIGNALS for c in top[0] if c is not None):
            ws = wb[name]
            break
    if ws is None:
        sys.exit(f"{path}: no sheet has the expected signal names in row 2 ({list(SIGNALS)}).")
    rows = ws.iter_rows(min_row=2, values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    idx = {SIGNALS[h]: i for i, h in enumerate(header) if h in SIGNALS}
    missing = set(SIGNALS.values()) - set(idx)
    if missing:
        sys.exit(f"{path}: missing signals for {sorted(missing)}")
    data = [[r[0]] + [r[idx[k]] for k in SIGNALS.values()] for r in rows if r and r[0] is not None]
    df = pd.DataFrame(data, columns=["time"] + list(SIGNALS.values()))
    df["time"] = pd.to_datetime(df["time"], errors="coerce")
    for k in SIGNALS.values():
        df[k] = pd.to_numeric(df[k], errors="coerce")
    return df


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", help="one or more .xlsx files followed by the output .csv")
    ap.add_argument("--basis", choices=["lifecycle", "direct"], default="lifecycle")
    ap.add_argument("--ef", action="append", default=[], metavar="SOURCE=VALUE")
    ap.add_argument("--shift-hours", type=float, default=0.0, help="add this to every timestamp (e.g. -5.5 for IST to UTC)")
    args = ap.parse_args()
    *xlsx, out = args.files
    if not xlsx or not out.lower().endswith(".csv"):
        sys.exit("Give one or more .xlsx files and then the output .csv path last.")
    ef = dict(FACTORS[args.basis])
    for s in args.ef:
        k, _, v = s.partition("=")
        if k not in ef:
            sys.exit(f"Unknown source {k!r}; use one of {list(ef)}")
        ef[k] = float(v)

    frames = []
    for p in xlsx:
        d = read_workbook(p)
        print(f"{p}: {len(d)} rows, {d['time'].min()} to {d['time'].max()}", flush=True)
        d["file"] = len(frames)
        frames.append(d)
    df = pd.concat(frames, ignore_index=True)
    n0 = len(df)
    srcs = list(SIGNALS.values())
    df = df.dropna()
    n_missing = n0 - len(df)
    # Solar reads a few MW below zero at night (sensor offset). That is not a bad sample, so clip it
    # to 0 instead of dropping the row; dropping would remove mostly night-time samples and bias the curve.
    n_clip = int((df["solar"] < 0).sum())
    df["solar"] = df["solar"].clip(lower=0)
    # Any other negative reading (e.g. gas at -3000 MW) or a zero coal reading is a sensor glitch: drop the row.
    keep = (df[srcs] >= 0).all(axis=1) & (df["coal"] > 0)
    n_glitch = int((~keep).sum())
    df = df[keep]
    tot = df[srcs].sum(axis=1)
    low = tot <= 0.5 * tot.median()              # total output far below normal: dropout
    n_low = int(low.sum())
    df = df[~low]
    print(f"Kept {len(df)} of {n0} samples. Missing: {n_missing}; sensor glitches dropped: {n_glitch}; "
          f"implausibly low total dropped: {n_low}; night-time solar clipped to 0 (rows kept): {n_clip}")

    df["total"] = df[srcs].sum(axis=1)
    df["emis"] = sum(df[k] * ef[k] for k in srcs)
    hour = (df["time"] + pd.Timedelta(hours=args.shift_hours)).dt.hour.rename("hour")
    # One curve per file (so a 10-second file does not outweigh a 5-minute file), then the plain
    # average of those curves: each month/year counts equally.
    g = df.groupby(["file", hour])
    per_file = g["emis"].sum() / g["total"].sum()                       # index: (file, hour)
    curve = per_file.groupby(level="hour").mean().round(1)
    shares = (g[srcs].sum().div(g["total"].sum(), axis=0) * 100).groupby(level="hour").mean()
    if len(curve) != 24:
        sys.exit(f"Only {len(curve)} hours have data.")
    pd.DataFrame({"hour": curve.index, "gco2_per_kwh": curve.values}).to_csv(out, index=False)
    mean_mw = g["total"].mean().groupby(level="hour").mean().round(0)
    diag = shares.round(1).assign(gco2_per_kwh=curve, mean_total_MW=mean_mw)
    diag_path = out[:-4] + "_mix.csv"
    diag.to_csv(diag_path, index_label="hour")

    print(f"\nFactors used ({args.basis}): {ef}")
    print(f"Wrote {out} and {diag_path}")
    print(f"Daily mean {curve.mean():.0f} g/kWh; lowest {curve.min():.0f} at hour {curve.idxmin()}, highest {curve.max():.0f} at hour {curve.idxmax()}")
    print(f"Swing (max-min)/mean = {(curve.max() - curve.min()) / curve.mean() * 100:.0f}%")
    print(f"Solar share peaks at hour {shares['solar'].idxmax()} ({shares['solar'].max():.1f}% of generation); "
          "if that is not around midday, the timestamps are not local time.")


if __name__ == "__main__":
    main()
